"""Export a validated mathematics worksheet, with no student answer leakage."""
import base64
from io import BytesIO


def export_math(req):
    from pathlib import Path
    spec = req['document']
    output = Path(req['output'])
    teacher = spec.get('teacher') is True
    image = None
    if spec.get('png'):
        image = base64.b64decode(spec['png'].split(',', 1)[1], validate=True)
        if not image.startswith(b'\x89PNG\r\n\x1a\n'):
            raise ValueError('Ảnh PNG không hợp lệ.')
    variants = spec.get('variants', [])
    footer = f"Phiên bản {spec['revision']} · Bản {'giáo viên' if teacher else 'học sinh'}"
    if output.suffix == '.docx':
        from docx import Document
        from docx.shared import Inches, Pt
        doc = Document()
        doc.styles['Normal'].font.name = 'Arial'
        doc.styles['Normal'].font.size = Pt(12)
        doc.add_heading(spec['title'], 0)
        doc.add_paragraph(footer)
        doc.add_paragraph(spec['description'])
        if image:
            doc.add_picture(BytesIO(image), width=Inches(5.5))
        doc.add_paragraph(spec['question'])
        for i, v in enumerate(variants, 1):
            doc.add_paragraph(f"Câu {i}. {v['question']}")
            doc.add_paragraph('Lời giải: ' + '.' * 65)
        if teacher:
            doc.add_page_break()
            doc.add_heading('Đáp án và hướng dẫn giáo viên', 1)
            for i, v in enumerate(variants, 1):
                doc.add_paragraph(f"Câu {i}. {v['answer']} — {v['check']}")
            if spec.get('answer'):
                doc.add_paragraph('Hướng dẫn tự nhập, chưa kiểm chứng: ' + spec['answer'])
            for check in spec['checks']:
                doc.add_paragraph(check)
        doc.save(output)
    elif output.suffix == '.pptx':
        from pptx import Presentation
        from pptx.util import Inches, Pt
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        def slide(title, body='', picture=False):
            s = prs.slides.add_slide(prs.slide_layouts[6])
            box = s.shapes.add_textbox(Inches(.5), Inches(.3), Inches(12.3), Inches(.9))
            box.text_frame.text = title
            box.text_frame.paragraphs[0].font.size = Pt(30)
            if picture and image:
                s.shapes.add_picture(BytesIO(image), Inches(.6), Inches(1.4), height=Inches(5.5))
            box = s.shapes.add_textbox(Inches(6.4 if picture else .8), Inches(1.6), Inches(6 if picture else 11.7), Inches(5.2))
            box.text_frame.word_wrap = True
            box.text_frame.text = body
            for p in box.text_frame.paragraphs:
                p.font.size = Pt(22)
            return s
        slide(spec['title'], footer + '\n' + spec['description'])
        slide('Dự đoán', spec['question'])
        if image:
            slide('Quan sát và giải thích', 'Điều gì thay đổi? Điều gì được giữ nguyên?\nHãy giải thích bằng lập luận.', True)
        for i, v in enumerate(variants, 1):
            slide(f"Câu {i}", v['question'])
        if teacher:
            for i, v in enumerate(variants, 1):
                slide(f"Đáp án câu {i}", v['answer'] + '\n' + v['check'])
            if spec.get('answer'):
                # Notes keep lengthy teacher guidance off the projected slide.
                s = slide('Hướng dẫn giáo viên', 'Hướng dẫn tự nhập trong ghi chú — chưa kiểm chứng.')
                s.notes_slide.notes_text_frame.text = spec['answer']
        prs.save(output)
    elif output.suffix == '.xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook()
        ws = wb.active
        ws.title = 'Học liệu'
        if variants:
            ws.append(['Câu hỏi'] + (['Đáp án', 'Kiểm tra'] if teacher else []))
            for v in variants:
                ws.append([v['question']] + ([v['answer'], v['check']] if teacher else []))
        else:
            ws.append(['x', 'y (gần đúng)'])
            for row in spec['rows']:
                ws.append(row)
        for row in ws:
            for c in row:
                if isinstance(c.value, str):
                    c.data_type = 's'
                c.alignment = Alignment(wrap_text=True, vertical='top')
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='315C50')
        ws.freeze_panes = 'A2'
        for col in ['A', 'B', 'C']:
            ws.column_dimensions[col].width = 45
        info = wb.create_sheet('Thông tin')
        for value in [spec['title'], footer, spec['description'], spec['question']]:
            info.append([value])
            info.cell(info.max_row, 1).data_type = 's'
        wb.save(output)
    else:
        raise ValueError('Định dạng xuất không hỗ trợ.')
    return {'path': str(output)}

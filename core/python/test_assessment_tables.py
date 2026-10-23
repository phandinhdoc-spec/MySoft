import unittest,tempfile
from pathlib import Path
from docx import Document
from documents import extract
class AssessmentTableTests(unittest.TestCase):
 def test_preserves_merged_headers_blank_zero_and_vertical_origin(self):
  with tempfile.TemporaryDirectory() as d:
   file=Path(d)/'matrix.docx';doc=Document();doc.add_paragraph('Ma trận')
   t=doc.add_table(rows=4,cols=3);t.cell(0,0).text='Chủ đề';t.cell(0,1).merge(t.cell(0,2)).text='Mức độ'
   t.cell(1,1).text='NB';t.cell(1,2).text='TH';t.cell(2,0).merge(t.cell(3,0)).text='Áp suất'
   t.cell(2,1).text='0';t.cell(2,2).text='2';t.cell(3,2).text='1';doc.add_paragraph('Hết');doc.save(file)
   segments=extract(file)['segments'];rows=[s for s in segments if s.get('section')=='Bảng 1']
   self.assertEqual(len(rows),4);self.assertIn('[R1C2–C3] Mức độ',rows[0]['text'])
   self.assertIn('[R3C2] 0',rows[2]['text']);self.assertIn('[R4C2] [trống]',rows[3]['text'])
   self.assertIn('gộp dọc từ R3C1',rows[3]['text']);self.assertEqual(segments[0]['text'],'Ma trận');self.assertEqual(segments[-1]['text'],'Hết')
 def test_retains_nested_tables_and_separate_table_numbers(self):
  with tempfile.TemporaryDirectory() as d:
   file=Path(d)/'nested.docx';doc=Document();outer=doc.add_table(rows=1,cols=1);outer.cell(0,0).text='Ngoài'
   inner=outer.cell(0,0).add_table(rows=1,cols=2);inner.cell(0,0).text='NB';inner.cell(0,1).text='3'
   doc.add_table(rows=1,cols=1).cell(0,0).text='Bảng sau';doc.save(file)
   rows=extract(file)['segments'];self.assertEqual([s['section'] for s in rows],['Bảng 1','Bảng 2','Bảng 3']);self.assertIn('[R1C2] 3',rows[1]['text'])
if __name__=='__main__':unittest.main()

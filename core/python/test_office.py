import unittest,tempfile,zipfile
from pathlib import Path
from office import office_job, blank_office
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

class OfficeTests(unittest.TestCase):
 def setUp(self):self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
 def tearDown(self):self.tmp.cleanup()
 def create(self,name,document):
  out=self.root/name;office_job({'op':'create_office','output':str(out),'document':document});return out
 def test_blank_launch_files_are_empty_and_do_not_overwrite(self):
  for name in ['new.docx', 'new.xlsx', 'new.pptx']:
   target=self.root/name
   blank_office({'output':str(target)})
   original=target.read_bytes()
   with self.assertRaises(FileExistsError):blank_office({'output':str(target)})
   self.assertEqual(target.read_bytes(),original)
  self.assertEqual([p.text for p in Document(self.root/'new.docx').paragraphs],[''])
  self.assertIsNone(load_workbook(self.root/'new.xlsx').active['A1'].value)
  slides=Presentation(self.root/'new.pptx').slides
  self.assertEqual(len(slides),1)
  self.assertEqual(len(slides[0].shapes),0)
 def test_word_structure_and_no_overwrite(self):
  spec={'title':'Đáp án KHTN','blocks':[{'type':'heading','text':'Phần I','level':1},{'type':'paragraph','text':'Nội dung tiếng Việt H₂O'},{'type':'table','rows':[['Câu','Đáp án'],[1,'A']]},{'type':'list','items':['Một','Hai']}]}
  out=self.create('test.docx',spec);doc=Document(out)
  self.assertEqual(doc.tables[0].cell(1,1).text,'A');self.assertTrue(any(p.style.name=='Heading 1' for p in doc.paragraphs))
  original=out.read_bytes()
  with self.assertRaises(ValueError):self.create('test.docx',spec)
  self.assertEqual(out.read_bytes(),original)
 def test_excel_formulas_chart_and_copy_edit(self):
  src=self.create('test.xlsx',{'sheets':[{'name':'Điểm','rows':[['Tên','Điểm'],['An',8],['Bình',9],['=literal',None]],'formulas':[{'cell':'B4','formula':'=AVERAGE(B2:B3)'}],'charts':[{'title':'Điểm','maxRow':3}]}]})
  wb=load_workbook(src);ws=wb['Điểm'];self.assertEqual(ws['B4'].value,'=AVERAGE(B2:B3)');self.assertEqual(ws['A4'].data_type,'s');self.assertEqual(len(ws._charts),1)
  original=src.read_bytes();out=self.root/'edited.xlsx'
  office_job({'op':'edit_excel','path':str(src),'output':str(out),'document':{'sheets':[{'name':'Điểm','cells':[{'cell':'B2','value':10}]}]}})
  self.assertEqual(load_workbook(out)['Điểm']['B2'].value,10);self.assertEqual(src.read_bytes(),original)
 def test_powerpoint_slides_tables_and_notes(self):
  out=self.create('test.pptx',{'slides':[{'title':'Bài học','bullets':['Ý một','Ý hai'],'notes':'Lời giảng'},{'title':'So sánh','rows':[['A','B'],[1,2]]}]})
  slides=Presentation(out).slides;self.assertEqual(len(slides),2);self.assertIn('Lời giảng',slides[0].notes_slide.notes_text_frame.text);self.assertTrue(any(s.has_table for s in slides[1].shapes))
 def test_visual_deck_embeds_images_shapes_and_honors_first_layout(self):
  spec={'theme':'sage','slides':[{'title':'Áp suất','layout':'cover','bullets':['Khám phá: Lực và diện tích'],'notes':'Lời giảng'},{'title':'Vận dụng','layout':'visual','bullets':['Thiết kế: Tăng diện tích tiếp xúc','Giải thích: Áp suất giảm khi diện tích tăng']}]}
  out=self.create('visual.pptx',spec);deck=Presentation(out)
  for slide in deck.slides:
   self.assertTrue(any(sh.shape_type==13 for sh in slide.shapes),'Slide phải có ảnh nhúng')
   self.assertTrue(any(sh.shape_type==1 for sh in slide.shapes),'Khối phải sửa được')
   for sh in slide.shapes:
    self.assertLessEqual(sh.left+sh.width,deck.slide_width+10)
    self.assertLessEqual(sh.top+sh.height,deck.slide_height+10)
  cards=Presentation(self.create('cards.pptx',{'slides':[{'title':'Nội dung','layout':'cards','bullets':['Một: Nội dung','Hai: Nội dung']}]}))
  self.assertFalse(any(sh.shape_type==13 for sh in cards.slides[0].shapes))
 def test_slide_embeds_supplied_image_without_external_links(self):
  import base64,io
  from PIL import Image
  img=io.BytesIO();Image.new('RGB',(120,80),'red').save(img,format='PNG')
  out=self.create('image.pptx',{'slides':[{'title':'Ảnh nguồn','layout':'visual','image':{'data':base64.b64encode(img.getvalue()).decode(),'caption':'Ảnh thí nghiệm'},'bullets':['Quan sát: Nêu kết quả']} ]})
  deck=Presentation(out);pictures=[sh for sh in deck.slides[0].shapes if sh.shape_type==13]
  self.assertEqual(len(pictures),1);self.assertEqual(pictures[0].image.blob,img.getvalue())
  with self.assertRaises(ValueError):self.create('invalid-image.pptx',{'slides':[{'title':'Ảnh','image':{'data':'bad'}}]})
 def test_rejects_external_formula_and_bad_grid_before_writing(self):
  for spec in [{'sheets':[{'name':'A','rows':[['a'],['b','c']]}]},{'sheets':[{'name':'A','rows':[['a']],'formulas':[{'cell':'B1','formula':'=WEBSERVICE("https://example.com")'}]}]}]:
   with self.assertRaises(ValueError):self.create('bad.xlsx',spec)
   self.assertFalse((self.root/'bad.xlsx').exists())
if __name__=='__main__':unittest.main()

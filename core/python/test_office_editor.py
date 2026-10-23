import unittest,tempfile,zipfile,hashlib
from pathlib import Path
from office import editor_job
from docx import Document
from openpyxl import Workbook,load_workbook
from pptx import Presentation
from pptx.util import Inches

class EditorTests(unittest.TestCase):
 def setUp(self):self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
 def tearDown(self):self.tmp.cleanup()
 def edit(self,p,label,value):
  before=p.read_bytes();view=editor_job({'op':'office_editor_read','path':str(p)})
  entry=next(e for g in view['groups'] for e in g['entries'] if e['label']==label)
  out=p.with_name('edited'+p.suffix)
  req={'op':'office_editor_save','path':str(p),'output':str(out),'hash':view['hash'],'edits':[{'id':entry['id'],'value':value}]}
  editor_job(req);self.assertEqual(p.read_bytes(),before)
  with zipfile.ZipFile(p) as a,zipfile.ZipFile(out) as b:
   self.assertEqual(a.namelist(),b.namelist())
   edited=entry['id'].split('#')[0]
   for name in a.namelist():
    if name not in [edited,'xl/workbook.xml']:self.assertEqual(a.read(name),b.read(name),name)
  with self.assertRaises(FileExistsError):editor_job(req)
  req['output']=str(p.with_name('conflict'+p.suffix));req['hash']='wrong'
  with self.assertRaisesRegex(ValueError,'thay đổi'):editor_job(req)
  self.assertFalse(Path(req['output']).exists());return out
 def test_excel(self):
  p=self.root/'test.xlsx';w=Workbook();s=w.active;s['A1']='Tên';s['B2']=12;s['C2']='=B2*2';s['B2'].number_format='0.00';w.save(p)
  out=self.edit(p,'B2','25.5');w=load_workbook(out);self.assertEqual(w.active['B2'].value,25.5);self.assertEqual(w.active['B2'].number_format,'0.00');self.assertEqual(w.active['C2'].value,'=B2*2')
 def test_word_runs_and_table(self):
  p=self.root/'test.docx';d=Document();r=d.add_paragraph().add_run('Xin chào');r.bold=True;d.add_table(rows=1,cols=1).cell(0,0).text='Giữ bảng';d.save(p)
  out=self.edit(p,'Đoạn chữ 1','Chữ mới');d=Document(out);self.assertEqual(d.paragraphs[0].text,'Chữ mới');self.assertTrue(d.paragraphs[0].runs[0].bold);self.assertEqual(d.tables[0].cell(0,0).text,'Giữ bảng')
 def test_powerpoint(self):
  p=self.root/'test.pptx';d=Presentation();slide=d.slides.add_slide(d.slide_layouts[6]);shape=slide.shapes.add_textbox(Inches(1),Inches(2),Inches(5),Inches(1));shape.text='Slide cũ';d.save(p)
  out=self.edit(p,'Đoạn chữ 1','Slide mới');s=Presentation(out).slides[0].shapes[0];self.assertEqual(s.text,'Slide mới');self.assertEqual(s.left,Inches(1));self.assertEqual(s.top,Inches(2))
if __name__=='__main__':unittest.main()

class WindowedEditorTests(unittest.TestCase):
 def test_word_windows_cover_every_run_once(self):
  with tempfile.TemporaryDirectory() as root:
   p=Path(root)/'long.docx';d=Document()
   for i in range(70):d.add_paragraph('Paragraph '+str(i))
   d.save(p);ids=[];first=editor_job({'op':'office_editor_read','path':str(p),'windowed':True})
   for page in range(first['pages']):
    view=editor_job({'op':'office_editor_read','path':str(p),'windowed':True,'page':page})
    entries=view['groups'][0]['entries'];self.assertLessEqual(len(entries),24);ids.extend(e['id'] for e in entries)
   self.assertEqual(len(ids),70);self.assertEqual(len(set(ids)),70)
 def test_excel_windows_bound_rows_and_columns(self):
  with tempfile.TemporaryDirectory() as root:
   p=Path(root)/'wide.xlsx';w=Workbook()
   for row in range(1,82):
    for col in range(1,35):w.active.cell(row,col,'Text')
   w.save(p);ids=set();first=editor_job({'op':'office_editor_read','path':str(p),'windowed':True})
   for page in range(first['pages']):
    for column in range(first['columnPages']):
     view=editor_job({'op':'office_editor_read','path':str(p),'windowed':True,'page':page,'column':column})
     entries=view['groups'][0]['entries'];self.assertLessEqual(len(entries),640);ids.update(e['id'] for e in entries)
   self.assertEqual(len(ids),81*34)
